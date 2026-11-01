import openpyxl
from django.core.files.uploadedfile import UploadedFile

from gge_marketplace.common.application.unit_of_work import UnitOfWork
from gge_marketplace.pipelines.infrastructure.file_storage import TaskFileStorage
from gge_marketplace.pipelines.models import PipelineTask, StagingOrderItem


class InitializePipelineUseCase:
    def __init__(self, uow: UnitOfWork, storage_service: TaskFileStorage):
        self.storage_service = storage_service
        self.uow = uow

    def execute(self, task_id: int, uploaded_file: UploadedFile) -> PipelineTask:
        with self.uow:
            task: PipelineTask = PipelineTask.objects.select_for_update().get(pk=task_id)

            file_storage_path = self.storage_service.save_task_file(
                task_id=task.id,
                uploaded_file=uploaded_file
            )

            try:
                task.initialize_with_file(
                    file_path=file_storage_path,
                    original_name=uploaded_file.name,
                    file_size_bytes=uploaded_file.size
                )
            except Exception:
                self.storage_service.delete_file(file_storage_path)
                raise

            task.save()

        # transaction.on_commit(lambda: run_etl_pipeline_task.delay(task_id=task.id))

        return task


class LoadOrdersFromXlsxToStagingUseCase:
    BATCH_SIZE = 5000  # Размер чанка для сохранения в БД
    MAX_PROGRESS_VALUE = 50  # Максимальное значение прогресса задачи. Остальное -- при импорте из staging в orders

    def __init__(self, uow: UnitOfWork, storage_service: TaskFileStorage):
        self.storage_service = storage_service
        self.uow = uow

    def execute(self, task_id: int) -> None:
        task: PipelineTask = PipelineTask.objects.get(pk=task_id)

        if task.state != PipelineTask.State.EXTRACTING:
            task.state = PipelineTask.State.EXTRACTING
            task.save(update_fields=['state'])

        file_path = task.params.get('file_path')

        if not file_path:
            raise ValueError(f"Отсутствует имя файла в задаче {task_id}")

        StagingOrderItem.objects.filter(task_id=task_id).delete()

        with self.storage_service.open_task_file(file_path) as f:
            wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
            sheet = wb.active

            if sheet is None:
                raise ValueError("В xlsx файле отсутствуют листы")

            # Параметр для расчёта прогресса
            total_rows = sheet.max_row if sheet.max_row else 100000

            rows_iterator = sheet.iter_rows(values_only=True)

            # Читаем и валидируем первую строку (заголовки)
            try:
                headers = next(rows_iterator)
            except StopIteration:
                raise ValueError("Файл пуст или поврежден")

            expected_headers = ['order_id', 'user_email', 'sku', 'quantity', 'price_on_buy']

            # Приводим к нижнему регистру и очищаем от пробелов для стабильности
            cleaned_headers = [str(h).strip().lower() for h in headers if h is not None]

            if not all(header in cleaned_headers for header in expected_headers):
                raise ValueError(f"Неверная структура файла. Ожидались заголовки: {expected_headers}")

            # Создаем карту индексов колонок, чтобы не зависеть от их порядка в файле
            col_map = {name: cleaned_headers.index(name) for name in expected_headers}

            # Построчно читаем данные чанками
            staging_pool = []
            processed_count = 0

            # Безопасно извлекаем данные по маппингу колонок
            # Если ячейка пустая, записываем пустую строку
            def get_val(field_name: str, current_row: tuple) -> str:
                idx = col_map[field_name]
                val = current_row[idx] if idx < len(current_row) else ''
                return str(val).strip() if val is not None else ''

            for row in rows_iterator:
                # Пропускаем абсолютно пустые строки
                if not any(row):
                    continue

                item = StagingOrderItem(
                    task=task,
                    order_external_id=get_val('order_id', row),
                    user_email=get_val('user_email', row),
                    product_sku=get_val('sku', row),
                    quantity=get_val('quantity', row),
                    price_on_buy=get_val('price_on_buy', row),
                    state=StagingOrderItem.State.NEW
                )
                staging_pool.append(item)
                processed_count += 1

                # Если набрался размер пакета, сбрасываем в базу
                if len(staging_pool) >= self.BATCH_SIZE:
                    self._flush_batch(staging_pool, task, processed_count, total_rows)
                    staging_pool = []

            # Сбрасываем остаток, если он есть
            if staging_pool:
                self._flush_batch(staging_pool, task, processed_count, total_rows)

            wb.close()

        task.state = PipelineTask.State.EXTRACTED
        task.progress = self.MAX_PROGRESS_VALUE
        task.save(update_fields=['state', 'progress'])
        self.storage_service.delete_file(file_path)

    def _flush_batch(self, items: list, task: PipelineTask, current_count: int, total: int) -> None:
        """Вспомогательный метод для атомарной записи пакета строк и обновления прогресса."""
        with self.uow:
            StagingOrderItem.objects.bulk_create(items)

            # Считаем приблизительный прогресс
            calculated_progress = int((current_count / total) * self.MAX_PROGRESS_VALUE)

            task.progress = calculated_progress
            task.save(update_fields=['progress'])
